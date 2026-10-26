from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import stat
import sys
import traceback
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils.datetime import from_excel


class ConversionError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class FileEvidence:
    absolute_path: str
    size_bytes: int
    sha256: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "absolutePath": self.absolute_path,
            "sizeBytes": self.size_bytes,
            "sha256": self.sha256,
        }


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def evidence(path: Path, reported_path: Path | None = None) -> FileEvidence:
    normalized = path.resolve(strict=True)
    return FileEvidence(
        absolute_path=str((reported_path or normalized).resolve()),
        size_bytes=normalized.stat().st_size,
        sha256=sha256_file(normalized),
    )


def json_bytes(value: Any, *, pretty: bool = False) -> bytes:
    if pretty:
        text = json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
            allow_nan=False,
        ) + "\n"
    else:
        text = json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        )
    return text.encode("utf-8")


def write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise ConversionError("OUTPUT_COLLISION", f"Refusing to overwrite existing output: {path}")
    path.write_bytes(data)


def write_json(path: Path, value: Any) -> None:
    write_bytes(path, json_bytes(value, pretty=True))


def load_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exception:
        raise ConversionError("JSON_READ_FAILED", f"Could not read JSON contract/evidence: {path.name}") from exception


def ensure_within_project(path: Path, project_root: Path, label: str) -> Path:
    resolved = path.resolve()
    root = project_root.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exception:
        raise ConversionError("OUTPUT_OUTSIDE_PROJECT", f"{label} must be inside the project workspace") from exception
    return resolved


def ensure_local_data_root(path: Path, project_root: Path) -> None:
    relative = path.resolve().relative_to(project_root.resolve())
    if not relative.parts or relative.parts[0] != ".local":
        raise ConversionError(
            "OUTPUT_NOT_IN_IGNORED_DATA_ROOT",
            "Manual history output must be under the project's .local directory",
        )


def validate_path_segment(value: Any, label: str) -> str:
    if not isinstance(value, str) or re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", value) is None:
        raise ConversionError("UNSAFE_OUTPUT_SEGMENT", f"{label} must be one safe filename segment")
    return value


def checked_output_path(path: Path, project_root: Path) -> Path:
    resolved = ensure_within_project(path, project_root, "Output path")
    ensure_local_data_root(resolved, project_root)
    return resolved


def parse_iso_date(value: str, label: str) -> dt.date:
    try:
        return dt.date.fromisoformat(value)
    except (TypeError, ValueError) as exception:
        raise ConversionError("CONTRACT_INVALID", f"Invalid ISO date in {label}") from exception


def parse_workbook_date(value: Any, epoch: dt.datetime) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, dt.datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    if isinstance(value, str):
        candidate = value.strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return dt.datetime.strptime(candidate, pattern).date()
            except ValueError:
                pass
    return None


def json_safe_source_value(value: Any, unknown: str) -> Any:
    if value is None or value == "":
        return unknown
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def normalize_numeric(value: Any, scale: int, unknown: str) -> int | float | str:
    if value is None or value == "":
        return unknown
    if isinstance(value, bool):
        raise ConversionError("METRIC_NOT_NUMERIC", "Boolean value found in numeric metric cell")
    if not isinstance(value, (int, float, Decimal)):
        raise ConversionError("METRIC_NOT_NUMERIC", "Non-numeric value found in mapped metric cell")
    try:
        number = Decimal(str(value)) * scale
    except InvalidOperation as exception:
        raise ConversionError("METRIC_NOT_NUMERIC", "Invalid decimal value found in metric cell") from exception
    if not number.is_finite():
        raise ConversionError("METRIC_NOT_FINITE", "Non-finite value found in metric cell")
    if number == number.to_integral_value():
        return int(number)
    return float(number.normalize())


def period_for_date(contract: dict[str, Any], stat_date: dt.date) -> dict[str, Any]:
    matches = []
    for period in contract["periods"]:
        start = parse_iso_date(period["validFrom"], "period.validFrom")
        end = parse_iso_date(period["validThrough"], "period.validThrough")
        if start <= stat_date <= end:
            matches.append(period)
    if len(matches) != 1:
        raise ConversionError("MAPPING_PERIOD_NOT_UNIQUE", "Statistic date does not resolve to exactly one mapping period")
    return matches[0]


def contract_column(value: str | None) -> int | None:
    return None if value is None else column_index_from_string(value) - 1


def validate_contract(contract: dict[str, Any]) -> None:
    required = {
        "schemaVersion",
        "mappingVersion",
        "sourceType",
        "projectRef",
        "worksheet",
        "unknownValue",
        "factSchemaVersion",
        "catalogVersion",
        "dimensions",
        "metricFields",
        "rateFields",
        "periods",
        "overlap",
    }
    if not required.issubset(contract):
        raise ConversionError("CONTRACT_INVALID", "Mapping contract is missing required keys")
    validate_path_segment(contract["mappingVersion"], "mappingVersion")
    metric_fields = contract["metricFields"]
    if len(metric_fields) != len(set(metric_fields)):
        raise ConversionError("CONTRACT_INVALID", "Metric field list contains duplicates")
    if not set(contract["rateFields"]).issubset(metric_fields):
        raise ConversionError("CONTRACT_INVALID", "Rate fields must be a subset of metric fields")
    for period in contract["periods"]:
        unknown_fields = set(period["availableMetrics"]) - set(metric_fields)
        if unknown_fields:
            raise ConversionError("CONTRACT_INVALID", "Period maps fields outside the metric catalog")


def exact_sentinel_map(contract: dict[str, Any], material_id: str, material_name: str) -> str:
    for rule in contract.get("exactSentinelRules", []):
        if material_id == rule["workbookMaterialId"] and material_name == rule["exactName"]:
            return rule["canonicalMaterialId"]
    sentinel_ids = {rule["workbookMaterialId"] for rule in contract.get("exactSentinelRules", [])}
    sentinel_names = {rule["exactName"] for rule in contract.get("exactSentinelRules", [])}
    if material_id in sentinel_ids or material_name in sentinel_names:
        raise ConversionError(
            "AIGC_SENTINEL_NOT_EXACT",
            "AIGC sentinel exception did not match both exact name and exact workbook ID",
        )
    return material_id


def convert_workbook_rows(
    source_copy: Path,
    contract: dict[str, Any],
) -> tuple[dict[dt.date, list[dict[str, Any]]], dict[str, Any]]:
    try:
        workbook = openpyxl.load_workbook(source_copy, read_only=True, data_only=True)
    except Exception as exception:
        raise ConversionError("WORKBOOK_OPEN_FAILED", "Workbook could not be opened read-only") from exception
    try:
        if contract["worksheet"] not in workbook.sheetnames:
            raise ConversionError("WORKSHEET_MISSING", "Required worksheet is missing")
        worksheet = workbook[contract["worksheet"]]
        rows_by_date: dict[dt.date, list[dict[str, Any]]] = defaultdict(list)
        ignored_row_count = 0
        selected_row_count = 0
        zero_spend_count = 0
        positive_spend_count = 0
        mapping_counts: Counter[str] = Counter()
        unknown_counts: Counter[str] = Counter()
        rate_fields = set(contract["rateFields"])
        metric_fields = contract["metricFields"]
        unknown = contract["unknownValue"]
        optional_columns = {
            item["field"]: contract_column(item["column"])
            for item in contract.get("sourceOptionalColumns", [])
        }

        for source_row_number, values_tuple in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            values = list(values_tuple)
            rank = values[1] if len(values) > 1 else None
            if isinstance(rank, bool) or not isinstance(rank, int):
                ignored_row_count += 1
                continue
            stat_date = parse_workbook_date(values[0] if values else None, workbook.epoch)
            material_name_value = values[2] if len(values) > 2 else None
            material_id_value = values[3] if len(values) > 3 else None
            if stat_date is None or material_id_value in (None, ""):
                raise ConversionError(
                    "SELECTED_ROW_REQUIRED_VALUE_MISSING",
                    "A row with an integer rank is missing date or material ID",
                )
            material_name = unknown if material_name_value in (None, "") else str(material_name_value)
            display_material_id = str(material_id_value)
            material_id = exact_sentinel_map(contract, display_material_id, material_name)
            period = period_for_date(contract, stat_date)
            duplicate_date_column = contract_column(period.get("duplicateDateColumn"))
            if duplicate_date_column is not None:
                duplicate_value = values[duplicate_date_column] if duplicate_date_column < len(values) else None
                duplicate_date = parse_workbook_date(duplicate_value, workbook.epoch)
                if duplicate_date != stat_date:
                    raise ConversionError(
                        "DUPLICATE_DATE_MISMATCH",
                        "The legacy duplicate date column does not match the primary statistic date",
                    )

            created_value = values[6] if len(values) > 6 else None
            row: dict[str, Any] = {
                "material_create_time_v2": json_safe_source_value(created_value, unknown),
                "material_id": material_id,
                "roi2_material_video_name": material_name,
                "roi2_material_video_type": unknown,
                "stat_time_day": stat_date.isoformat(),
            }
            available = period["availableMetrics"]
            for field in metric_fields:
                column = available.get(field)
                if column is None:
                    row[field] = unknown
                    unknown_counts[field] += 1
                    continue
                index = contract_column(column)
                value = values[index] if index is not None and index < len(values) else None
                normalized = normalize_numeric(value, 100 if field in rate_fields else 1, unknown)
                row[field] = normalized
                if normalized == unknown:
                    unknown_counts[field] += 1

            spend = row["stat_cost_for_roi2"]
            if not isinstance(spend, (int, float)) or isinstance(spend, bool):
                raise ConversionError("SPEND_REQUIRED", "Every selected row must have a numeric spend value")
            if spend == 0:
                zero_spend_count += 1
            elif spend > 0:
                positive_spend_count += 1
            else:
                raise ConversionError("NEGATIVE_SPEND", "Negative spend is not accepted by this mapping")

            extensions = {
                field: json_safe_source_value(values[index] if index < len(values) else None, unknown)
                for field, index in optional_columns.items()
            }
            row["_manual_history"] = {
                "display_material_id": display_material_id,
                "mapping_period": period["id"],
                "source_optional_columns": extensions,
                "source_rank": rank,
                "source_workbook_row": source_row_number,
            }
            rows_by_date[stat_date].append(row)
            selected_row_count += 1
            mapping_counts[period["id"]] += 1

        if not rows_by_date:
            raise ConversionError("NO_FACT_ROWS", "Workbook contains no selected material-day rows")

        rank_anomalies: dict[str, dict[str, int]] = {}
        for stat_date, rows in rows_by_date.items():
            ranks = [row["_manual_history"]["source_rank"] for row in rows]
            if sorted(ranks) != list(range(1, len(rows) + 1)):
                rank_counts = Counter(ranks)
                rank_anomalies[stat_date.isoformat()] = {
                    "rowCount": len(ranks),
                    "minimumRank": min(ranks),
                    "maximumRank": max(ranks),
                    "missingWithinMinMax": len(set(range(min(ranks), max(ranks) + 1)) - set(ranks)),
                    "duplicateRankCount": sum(count - 1 for count in rank_counts.values() if count > 1),
                }
            rows.sort(key=lambda item: (
                item["_manual_history"]["source_rank"],
                item["_manual_history"]["source_workbook_row"],
            ))
            material_ids = [row["material_id"] for row in rows]
            if len(material_ids) != len(set(material_ids)):
                raise ConversionError("DUPLICATE_MATERIAL_DAY_KEY", f"Duplicate material ID for {stat_date.isoformat()}")

        stats = {
            "worksheetPhysicalRowsExcludingHeader": worksheet.max_row - 1,
            "selectedFactRows": selected_row_count,
            "ignoredRows": ignored_row_count,
            "datesWithRows": len(rows_by_date),
            "firstDate": min(rows_by_date).isoformat(),
            "lastDate": max(rows_by_date).isoformat(),
            "zeroSpendRows": zero_spend_count,
            "positiveSpendRows": positive_spend_count,
            "mappingPeriodRowCounts": dict(sorted(mapping_counts.items())),
            "unknownValueCountsByField": dict(sorted(unknown_counts.items())),
            "sourceRankAnomalyDays": rank_anomalies,
        }
        return dict(rows_by_date), stats
    finally:
        workbook.close()


def numbers_equal(left: Any, right: Any) -> bool:
    if isinstance(left, bool) or isinstance(right, bool):
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(Decimal(str(left)) - Decimal(str(right))) <= Decimal("0.000000001")
    return left == right


def overlap_file_paths(auto_root: Path, project_ref: str, stat_date: dt.date) -> tuple[Path, Path]:
    directory = auto_root / stat_date.isoformat() / project_ref
    return directory / "manifest.json", directory / "material_day.json"


def compare_overlap_day(
    auto_root: Path,
    project_ref: str,
    stat_date: dt.date,
    manual_rows: list[dict[str, Any]],
    compare_fields: list[str],
) -> tuple[dict[str, Any], str | None]:
    manifest_path, material_path = overlap_file_paths(auto_root, project_ref, stat_date)
    if not manifest_path.is_file() or not material_path.is_file():
        return {
            "status": "AUTO_EVIDENCE_MISSING",
            "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE",
        }, None
    manifest = load_json(manifest_path)
    material = load_json(material_path)
    if (
        manifest.get("jobStatus") != "COMPLETE"
        or manifest.get("paginationComplete") is not True
        or material.get("paginationComplete") is not True
        or manifest.get("statDate") != stat_date.isoformat()
        or material.get("statDate") != stat_date.isoformat()
        or manifest.get("projectRef") != project_ref
        or material.get("projectRef") != project_ref
        or manifest.get("advertiserId") != material.get("advertiserId")
    ):
        return {
            "status": "AUTO_EVIDENCE_NOT_COMPLETE",
            "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE",
        }, None

    auto_rows = material.get("rows")
    if not isinstance(auto_rows, list):
        return {
            "status": "AUTO_EVIDENCE_INVALID",
            "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE",
        }, None
    positive_auto = [
        row for row in auto_rows
        if isinstance(row, dict)
        and isinstance(row.get("stat_cost_for_roi2"), (int, float))
        and not isinstance(row.get("stat_cost_for_roi2"), bool)
        and row["stat_cost_for_roi2"] > 0
    ]
    auto_by_id = {str(row.get("material_id")): row for row in positive_auto}
    manual_by_id = {str(row.get("material_id")): row for row in manual_rows}
    missing_in_manual = len(set(auto_by_id) - set(manual_by_id))
    extra_in_manual = len(set(manual_by_id) - set(auto_by_id))
    field_mismatches: Counter[str] = Counter()
    compared_rows = 0
    for material_id in sorted(set(auto_by_id) & set(manual_by_id)):
        compared_rows += 1
        auto_row = auto_by_id[material_id]
        manual_row = manual_by_id[material_id]
        for field in compare_fields:
            if not numbers_equal(manual_row.get(field), auto_row.get(field)):
                field_mismatches[field] += 1

    zero_count = sum(
        1 for row in auto_rows
        if isinstance(row, dict)
        and isinstance(row.get("stat_cost_for_roi2"), (int, float))
        and not isinstance(row.get("stat_cost_for_roi2"), bool)
        and row["stat_cost_for_roi2"] == 0
    )
    matched = (
        len(manual_rows) == len(positive_auto)
        and missing_in_manual == 0
        and extra_in_manual == 0
        and not field_mismatches
    )
    result = {
        "status": "MATCHED_AUTO_POSITIVE_ROWS" if matched else "MISMATCH_BLOCKED",
        "authoritativeSource": "MCP_AUTO",
        "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE",
        "manualRows": len(manual_rows),
        "autoRows": len(auto_rows),
        "autoPositiveRows": len(positive_auto),
        "autoZeroSpendRows": zero_count,
        "comparedRows": compared_rows,
        "comparedFieldCount": len(compare_fields),
        "fieldComparisons": compared_rows * len(compare_fields),
        "missingInManualCount": missing_in_manual,
        "extraInManualCount": extra_in_manual,
        "fieldMismatchCounts": dict(sorted(field_mismatches.items())),
        "manifestEvidence": {
            "sizeBytes": manifest_path.stat().st_size,
            "sha256": sha256_file(manifest_path),
        },
        "materialEvidence": {
            "sizeBytes": material_path.stat().st_size,
            "sha256": sha256_file(material_path),
        },
    }
    return result, str(material.get("advertiserId"))


def validate_attestation_binding(
    payload: dict[str, Any], schema: str, account_id: str | None, project_ref: str, source_hash: str
) -> None:
    if (
        not isinstance(payload, dict)
        or payload.get("schemaVersion") != schema
        or payload.get("status") != "VERIFIED"
        or payload.get("verificationMethod") != "USER_CONFIRMED"
        or not isinstance(payload.get("evidenceRef"), str)
        or not payload["evidenceRef"].strip()
        or not account_id
        or payload.get("accountId") != account_id
        or payload.get("projectRef") != project_ref
        or payload.get("sourceWorkbookSha256") != source_hash
    ):
        raise ConversionError("ATTESTATION_BINDING_INVALID", "Attestation must bind the verified account, project and exact source workbook")


def load_identity_binding(
    account_id: str | None, evidence_path: Path | None, project_ref: str, source_hash: str
) -> dict[str, Any]:
    if account_id is None and evidence_path is None:
        return {"status": "MISSING", "accountId": None, "evidence": None}
    if not account_id or evidence_path is None:
        raise ConversionError("ACCOUNT_BINDING_INCOMPLETE", "Account ID and account evidence must be supplied together")
    payload = load_json(evidence_path)
    validate_attestation_binding(
        payload, "manual-history-account-attestation-1.0.0", account_id, project_ref, source_hash
    )
    valid_from = parse_iso_date(payload.get("validFrom"), "identity.validFrom")
    valid_through = parse_iso_date(payload.get("validThrough"), "identity.validThrough")
    if valid_from > valid_through:
        raise ConversionError("ACCOUNT_BINDING_INVALID", "Account identity date range is reversed")
    return {
        "status": "VERIFIED",
        "accountId": account_id,
        "evidence": evidence(evidence_path).as_dict(),
        "validFrom": valid_from.isoformat(),
        "validThrough": valid_through.isoformat(),
    }


def load_completeness_binding(
    path: Path | None, account_id: str | None, project_ref: str, source_hash: str
) -> tuple[dict[str, str], dict[str, Any] | None]:
    if path is None:
        return {}, None
    payload = load_json(path)
    validate_attestation_binding(
        payload, "manual-history-completeness-attestation-1.0.0", account_id, project_ref, source_hash
    )
    days = payload.get("days")
    if payload.get("status") != "VERIFIED" or not isinstance(days, dict):
        raise ConversionError("COMPLETENESS_EVIDENCE_INVALID", "Completeness evidence must be VERIFIED and contain a days map")
    allowed = {"COMPLETE", "COMPLETE_NO_ROWS", "UNVERIFIED"}
    normalized: dict[str, str] = {}
    for key, value in days.items():
        parse_iso_date(key, "completeness.days key")
        if value not in allowed:
            raise ConversionError("COMPLETENESS_EVIDENCE_INVALID", "Completeness day status is unsupported")
        normalized[key] = value
    return normalized, evidence(path).as_dict()


def copy_source_evidence(source: Path, output_root: Path, source_hash: str, project_root: Path) -> tuple[Path, FileEvidence]:
    destination_dir = checked_output_path(output_root / "sources" / source_hash, project_root)
    destination = destination_dir / "source.xlsx"
    destination_dir.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        copied = evidence(destination)
        if copied.sha256 != source_hash or copied.size_bytes != source.stat().st_size:
            raise ConversionError("SOURCE_EVIDENCE_COLLISION", "Existing source evidence does not match the requested source")
        return destination, copied
    temporary = destination_dir / "source.xlsx.copying"
    if temporary.exists():
        raise ConversionError("SOURCE_COPY_INCOMPLETE", "An incomplete source evidence copy already exists")
    shutil.copyfile(source, temporary)
    copied_hash = sha256_file(temporary)
    if copied_hash != source_hash or temporary.stat().st_size != source.stat().st_size:
        raise ConversionError("SOURCE_COPY_MISMATCH", "Copied source evidence failed byte-for-byte verification")
    os.replace(temporary, destination)
    destination.chmod(stat.S_IREAD)
    return destination, evidence(destination)


def daily_content_hash(rows: list[dict[str, Any]]) -> str:
    return sha256_bytes(json_bytes(rows))


def facts_file_bytes(rows: list[dict[str, Any]]) -> bytes:
    if not rows:
        return b""
    return b"".join(json_bytes(row) + b"\n" for row in rows)


def date_range(first: dt.date, last: dt.date) -> Iterable[dt.date]:
    current = first
    while current <= last:
        yield current
        current += dt.timedelta(days=1)


def tree_evidence(root: Path, excluded_names: set[str] | None = None) -> dict[str, Any]:
    excluded_names = excluded_names or set()
    files = []
    total_size = 0
    for path in sorted(item for item in root.rglob("*") if item.is_file() and item.name not in excluded_names):
        relative = path.relative_to(root).as_posix()
        size = path.stat().st_size
        digest = sha256_file(path)
        files.append({"path": relative, "sizeBytes": size, "sha256": digest})
        total_size += size
    canonical = "".join(
        f"{item['path']}\t{item['sizeBytes']}\t{item['sha256']}\n" for item in files
    ).encode("utf-8")
    return {
        "fileCount": len(files),
        "sizeBytes": total_size,
        "sha256": sha256_bytes(canonical),
        "files": files,
    }


def readback_bundle(stage: Path, expected_days: list[dict[str, Any]]) -> dict[str, Any]:
    total_rows = 0
    zero_rows = 0
    dates = 0
    for expected in expected_days:
        stat_date = expected["statDate"]
        day_dir = stage / "days" / stat_date
        day_evidence_path = day_dir / "day-evidence.json"
        facts_path = day_dir / "facts.jsonl"
        day_evidence = load_json(day_evidence_path)
        facts_bytes = facts_path.read_bytes()
        if sha256_bytes(facts_bytes) != day_evidence["factsFileEvidence"]["sha256"]:
            raise ConversionError("READBACK_FILE_HASH_MISMATCH", f"Facts file hash mismatch for {stat_date}")
        rows = []
        for line_number, line in enumerate(facts_bytes.splitlines(), start=1):
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exception:
                raise ConversionError("READBACK_JSON_INVALID", f"Invalid JSONL at {stat_date}:{line_number}") from exception
            if row.get("stat_time_day") != stat_date:
                raise ConversionError("READBACK_DATE_MISMATCH", f"Row date mismatch for {stat_date}")
            rows.append(row)
        if len(rows) != day_evidence["rowCount"]:
            raise ConversionError("READBACK_ROW_COUNT_MISMATCH", f"Row count mismatch for {stat_date}")
        if daily_content_hash(rows) != day_evidence["contentHash"]:
            raise ConversionError("READBACK_CONTENT_HASH_MISMATCH", f"Content hash mismatch for {stat_date}")
        ids = [str(row.get("material_id")) for row in rows]
        if len(ids) != len(set(ids)):
            raise ConversionError("READBACK_DUPLICATE_KEY", f"Duplicate material ID for {stat_date}")
        actual_zero = sum(1 for row in rows if row.get("stat_cost_for_roi2") == 0)
        if actual_zero != day_evidence["zeroSpendRowCount"]:
            raise ConversionError("READBACK_ZERO_COUNT_MISMATCH", f"Zero-spend count mismatch for {stat_date}")
        total_rows += len(rows)
        zero_rows += actual_zero
        dates += 1
    return {
        "status": "PASSED",
        "daysReadBack": dates,
        "rowsReadBack": total_rows,
        "zeroSpendRowsReadBack": zero_rows,
        "method": "Every JSONL line reparsed; dates, unique keys, row counts, zero counts, file SHA-256 and canonical content SHA-256 recomputed",
    }


def make_binding_fingerprint(
    identity: dict[str, Any],
    completeness_evidence: dict[str, Any] | None,
) -> str:
    payload = {
        "accountId": identity.get("accountId"),
        "identityEvidenceSha256": (identity.get("evidence") or {}).get("sha256"),
        "completenessEvidenceSha256": (completeness_evidence or {}).get("sha256"),
    }
    return sha256_bytes(json_bytes(payload))[:16]


def set_tree_read_only(root: Path) -> None:
    for path in root.rglob("*"):
        if path.is_file():
            path.chmod(stat.S_IREAD)


def convert(args: argparse.Namespace) -> dict[str, Any]:
    project_root = Path(args.project_root).resolve(strict=True)
    validate_path_segment(args.run_id, "run_id")
    output_root = ensure_within_project(Path(args.output_root), project_root, "Output root")
    ensure_local_data_root(output_root, project_root)
    contract_path = ensure_within_project(Path(args.contract), project_root, "Mapping contract")
    source = Path(args.source).resolve(strict=True)
    if not source.is_file():
        raise ConversionError("SOURCE_NOT_FILE", "Source workbook is not a regular file")
    contract = load_json(contract_path)
    validate_contract(contract)
    if args.project_ref != contract["projectRef"]:
        raise ConversionError("PROJECT_REF_MISMATCH", "Requested project does not match the mapping contract")

    source_size = source.stat().st_size
    source_hash = sha256_file(source)
    if source_hash.lower() != args.expected_source_sha256.lower():
        raise ConversionError("SOURCE_HASH_MISMATCH", "Source workbook SHA-256 does not match the expected value")
    output_root.mkdir(parents=True, exist_ok=True)
    usage = shutil.disk_usage(output_root)
    estimated_required = source_size * 8 + 64 * 1024 * 1024
    if usage.free < estimated_required:
        raise ConversionError("INSUFFICIENT_CAPACITY", "Output volume does not have the required preflight capacity")

    source_copy, source_copy_evidence = copy_source_evidence(source, output_root, source_hash, project_root)
    identity = load_identity_binding(
        args.account_id,
        Path(args.account_evidence).resolve(strict=True) if args.account_evidence else None,
        args.project_ref,
        source_hash,
    )
    completeness, completeness_evidence = load_completeness_binding(
        Path(args.completeness_evidence).resolve(strict=True) if args.completeness_evidence else None,
        args.account_id, args.project_ref, source_hash,
    )
    binding_fingerprint = make_binding_fingerprint(identity, completeness_evidence)
    bundle_id = f"{source_hash[:16]}-{contract['mappingVersion']}-{binding_fingerprint}"
    final = checked_output_path(
        output_root / "derived" / source_hash / contract["mappingVersion"] / binding_fingerprint,
        project_root,
    )
    if final.exists():
        receipt_path = final / "conversion-receipt.json"
        if not receipt_path.is_file():
            raise ConversionError("OUTPUT_COLLISION", "Final bundle directory exists without a receipt")
        receipt = load_json(receipt_path)
        current_contract_hash = sha256_file(contract_path)
        if (
            receipt.get("bundleId") != bundle_id
            or receipt.get("sourceWorkbook", {}).get("sha256") != source_hash
            or receipt.get("mappingContractEvidence", {}).get("sha256") != current_contract_hash
        ):
            raise ConversionError("OUTPUT_COLLISION", "Existing final bundle identity does not match")
        recorded_tree = receipt.get("output", {})
        actual_tree = tree_evidence(final, {"conversion-receipt.json"})
        if (
            recorded_tree.get("fileCount") != actual_tree["fileCount"]
            or recorded_tree.get("sizeBytes") != actual_tree["sizeBytes"]
            or recorded_tree.get("treeSha256") != actual_tree["sha256"]
        ):
            raise ConversionError("EXISTING_BUNDLE_READBACK_FAILED", "Existing final bundle no longer matches its receipt")
        return receipt

    staging = checked_output_path(output_root / "staging" / args.run_id, project_root)
    if staging.exists():
        raise ConversionError("STAGING_COLLISION", "Run staging directory already exists")
    staging.mkdir(parents=True)

    rows_by_date, workbook_stats = convert_workbook_rows(source_copy, contract)
    if sha256_file(source) != source_hash or source.stat().st_size != source_size:
        raise ConversionError("SOURCE_CHANGED_DURING_CONVERSION", "Desktop source changed during conversion")
    if sha256_file(source_copy) != source_hash:
        raise ConversionError("SOURCE_EVIDENCE_CHANGED", "Project source evidence changed during conversion")

    overlap_root = Path(args.auto_overlap_root).resolve(strict=True) if args.auto_overlap_root else None
    overlap_dates = {parse_iso_date(value, "overlap date") for value in contract["overlap"]["dates"]}
    overlap_results: dict[str, dict[str, Any]] = {}
    observed_auto_accounts: set[str] = set()
    if overlap_root:
        for stat_date in sorted(overlap_dates):
            result, account = compare_overlap_day(
                overlap_root,
                args.project_ref,
                stat_date,
                rows_by_date.get(stat_date, []),
                contract["overlap"]["compareFields"],
            )
            overlap_results[stat_date.isoformat()] = result
            if account:
                observed_auto_accounts.add(account)
        failures = [date for date, item in overlap_results.items() if item["status"] != "MATCHED_AUTO_POSITIVE_ROWS"]
        if failures:
            raise ConversionError("OVERLAP_COMPARISON_FAILED", "One or more required overlap days did not match automatic facts")
        if len(observed_auto_accounts) != 1:
            raise ConversionError("OVERLAP_ACCOUNT_INCONSISTENT", "Automatic overlap evidence does not resolve to one account")

    first_date = min(rows_by_date)
    last_date = max(rows_by_date)
    day_summaries: list[dict[str, Any]] = []
    day_manifest_records: list[dict[str, Any]] = []
    for stat_date in date_range(first_date, last_date):
        rows = rows_by_date.get(stat_date, [])
        day_dir = staging / "days" / stat_date.isoformat()
        facts_path = day_dir / "facts.jsonl"
        facts_bytes = facts_file_bytes(rows)
        write_bytes(facts_path, facts_bytes)
        final_facts_path = final / "days" / stat_date.isoformat() / "facts.jsonl"
        facts_evidence = evidence(facts_path, reported_path=final_facts_path)
        completeness_status = completeness.get(
            stat_date.isoformat(),
            "UNVERIFIED_NO_ROWS" if not rows else "UNVERIFIED",
        )
        if (completeness_status == "COMPLETE_NO_ROWS" and rows) or (completeness_status == "COMPLETE" and not rows):
            raise ConversionError("COMPLETENESS_ROW_COUNT_MISMATCH", "Attested day completeness contradicts the source row count")
        overlap = overlap_results.get(
            stat_date.isoformat(),
            {
                "status": "AUTO_EVIDENCE_NOT_PROVIDED" if stat_date in overlap_dates else "NOT_IN_OVERLAP_SCOPE",
                "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE" if stat_date in overlap_dates else "BLOCK_UNTIL_VERIFIED",
            },
        )
        blockers = []
        if identity["status"] != "VERIFIED":
            blockers.append("MISSING_ACCOUNT_IDENTITY_EVIDENCE")
        elif not identity["validFrom"] <= stat_date.isoformat() <= identity["validThrough"]:
            blockers.append("ACCOUNT_IDENTITY_DATE_OUTSIDE_ATTESTED_RANGE")
        if completeness_status not in {"COMPLETE", "COMPLETE_NO_ROWS"}:
            blockers.append("MISSING_DAY_COMPLETENESS_EVIDENCE")
        if overlap["status"] != "NOT_IN_OVERLAP_SCOPE":
            blockers.append("AUTO_AUTHORITATIVE_OVERLAP_DO_NOT_IMPORT")
        database_import_eligible = not blockers
        status = "READY_FOR_IMPORT" if database_import_eligible else "STAGED_UNVERIFIED"
        zero_count = sum(1 for row in rows if row["stat_cost_for_roi2"] == 0)
        unknown_count = sum(
            1
            for row in rows
            for field in contract["metricFields"]
            if row[field] == contract["unknownValue"]
        )
        content_hash = daily_content_hash(rows)
        source_ranks = [row["_manual_history"]["source_rank"] for row in rows]
        source_rank_status = (
            "NO_ROWS"
            if not rows
            else "CONTIGUOUS_1_TO_N"
            if sorted(source_ranks) == list(range(1, len(rows) + 1))
            else "SPARSE_OR_DUPLICATE_PRESERVED"
        )
        data_version = (
            f"manual-history-{stat_date.isoformat()}-"
            f"{contract['mappingVersion']}-{source_hash[:12]}-{content_hash[:12]}"
        )
        day_evidence = {
            "schemaVersion": "manual-history-day-evidence-1.0.0",
            "sourceType": "MANUAL_HISTORY",
            "runId": args.run_id,
            "projectRef": args.project_ref,
            "accountId": identity["accountId"],
            "accountIdentityStatus": identity["status"],
            "accountIdentityEvidence": identity["evidence"],
            "statDate": stat_date.isoformat(),
            "status": status,
            "databaseImportEligible": database_import_eligible,
            "publicationBlockers": blockers,
            "completenessStatus": completeness_status,
            "completenessEvidence": completeness_evidence,
            "mappingVersion": contract["mappingVersion"],
            "factSchemaVersion": contract["factSchemaVersion"],
            "catalogVersion": contract["catalogVersion"],
            "dataVersion": data_version,
            "contentHash": content_hash,
            "rowCount": len(rows),
            "zeroSpendRowCount": zero_count,
            "unknownMetricValueCount": unknown_count,
            "sourceRankStatus": source_rank_status,
            "sourceWorkbookEvidence": source_copy_evidence.as_dict(),
            "factsFileEvidence": facts_evidence.as_dict(),
            "mcpManifestStatus": "NOT_APPLICABLE_NOT_CREATED",
            "overlap": overlap,
        }
        write_json(day_dir / "day-evidence.json", day_evidence)
        summary = {
            "statDate": stat_date.isoformat(),
            "status": status,
            "databaseImportEligible": database_import_eligible,
            "completenessStatus": completeness_status,
            "rowCount": len(rows),
            "zeroSpendRowCount": zero_count,
            "unknownMetricValueCount": unknown_count,
            "sourceRankStatus": source_rank_status,
            "mappingVersion": contract["mappingVersion"],
            "contentHash": content_hash,
            "factsSizeBytes": facts_evidence.size_bytes,
            "factsSha256": facts_evidence.sha256,
            "overlapStatus": overlap["status"],
            "publicationBlockers": blockers,
        }
        day_summaries.append(summary)
        day_manifest_records.append({
            "statDate": stat_date.isoformat(),
            "evidencePath": f"days/{stat_date.isoformat()}/day-evidence.json",
            "factsPath": f"days/{stat_date.isoformat()}/facts.jsonl",
            "rowCount": len(rows),
            "contentHash": content_hash,
        })

    daily_summary_bytes = b"".join(json_bytes(item) + b"\n" for item in day_summaries)
    write_bytes(staging / "daily-summary.jsonl", daily_summary_bytes)
    bundle_manifest = {
        "schemaVersion": "manual-history-bundle-manifest-1.0.0",
        "bundleId": bundle_id,
        "runId": args.run_id,
        "sourceType": "MANUAL_HISTORY",
        "mappingVersion": contract["mappingVersion"],
        "projectRef": args.project_ref,
        "accountIdentityStatus": identity["status"],
        "sourceWorkbookSha256": source_hash,
        "days": day_manifest_records,
    }
    write_json(staging / "bundle-manifest.json", bundle_manifest)
    readback = readback_bundle(staging, day_summaries)
    output_tree = tree_evidence(staging)
    evidence_index = {
        "schemaVersion": "manual-history-evidence-index-1.0.0",
        "scope": "All generated files before conversion-receipt.json",
        "fileCount": output_tree["fileCount"],
        "sizeBytes": output_tree["sizeBytes"],
        "treeSha256": output_tree["sha256"],
        "files": output_tree["files"],
    }
    write_json(staging / "evidence-index.json", evidence_index)
    final_tree = tree_evidence(staging)

    staged_days = sum(1 for item in day_summaries if item["status"] == "STAGED_UNVERIFIED")
    ready_days = len(day_summaries) - staged_days
    overlap_summary = {
        "requiredDays": len(overlap_dates),
        "comparedDays": len(overlap_results),
        "matchedDays": sum(1 for item in overlap_results.values() if item["status"] == "MATCHED_AUTO_POSITIVE_ROWS"),
        "manualRows": sum(item.get("manualRows", 0) for item in overlap_results.values()),
        "autoPositiveRows": sum(item.get("autoPositiveRows", 0) for item in overlap_results.values()),
        "autoZeroSpendRows": sum(item.get("autoZeroSpendRows", 0) for item in overlap_results.values()),
        "fieldComparisons": sum(item.get("fieldComparisons", 0) for item in overlap_results.values()),
        "fieldMismatchCount": sum(sum(item.get("fieldMismatchCounts", {}).values()) for item in overlap_results.values()),
        "ordinaryKeyMissingOrExtraCount": sum(
            item.get("missingInManualCount", 0) + item.get("extraInManualCount", 0)
            for item in overlap_results.values()
        ),
        "manualPublicationAction": "DO_NOT_IMPORT_OR_OVERWRITE",
    }
    receipt = {
        "schemaVersion": "manual-history-conversion-receipt-1.0.0",
        "bundleId": bundle_id,
        "runId": args.run_id,
        "result": "STAGED_UNVERIFIED" if staged_days else "READY_FOR_IMPORT",
        "sourceType": "MANUAL_HISTORY",
        "projectRef": args.project_ref,
        "mappingVersion": contract["mappingVersion"],
        "mappingContractEvidence": evidence(contract_path).as_dict(),
        "sourceWorkbook": {
            "originalPath": str(source),
            "sizeBytes": source_size,
            "sha256": source_hash,
            "lastWriteTimeUtc": dt.datetime.fromtimestamp(source.stat().st_mtime, dt.timezone.utc).isoformat(),
        },
        "sourceEvidenceCopy": source_copy_evidence.as_dict(),
        "capacityPreflight": {
            "volumeFreeBytes": usage.free,
            "estimatedRequiredBytes": estimated_required,
            "passed": True,
        },
        "accountIdentity": identity,
        "workbook": workbook_stats,
        "daily": {
            "calendarDays": len(day_summaries),
            "daysWithRows": sum(1 for item in day_summaries if item["rowCount"] > 0),
            "daysWithoutRows": sum(1 for item in day_summaries if item["rowCount"] == 0),
            "readyForImportDays": ready_days,
            "stagedUnverifiedDays": staged_days,
        },
        "overlap": overlap_summary,
        "readback": readback,
        "output": {
            "absolutePath": str(final),
            "treeScope": "All bundle files except conversion-receipt.json",
            "fileCount": final_tree["fileCount"],
            "sizeBytes": final_tree["sizeBytes"],
            "treeSha256": final_tree["sha256"],
        },
        "database": {
            "importAttempted": False,
            "readbackStatus": "NOT_ATTEMPTED_NO_VERIFIED_ACCOUNT_DAY_OR_MYSQL_ENVIRONMENT",
        },
        "failureIsolation": {
            "stagingPath": str(staging),
            "publicationIsAtomic": True,
            "rule": "Only a fully written and fully read-back bundle is renamed into the derived evidence path",
        },
        "limitations": [
            "No MCP manifest was created or implied for the manual workbook.",
            "No date is publishable without independent account identity and day-completeness evidence.",
            "Automatic overlap days remain authoritative and manual rows are blocked from import or overwrite.",
            "No MySQL connection or database import was attempted.",
        ],
    }
    write_json(staging / "conversion-receipt.json", receipt)
    final.parent.mkdir(parents=True, exist_ok=True)
    os.replace(staging, final)
    set_tree_read_only(final)
    return receipt


def write_failure_receipt(args: argparse.Namespace, error: ConversionError) -> None:
    try:
        project_root = Path(args.project_root).resolve(strict=True)
        validate_path_segment(args.run_id, "run_id")
        output_root = ensure_within_project(Path(args.output_root), project_root, "Output root")
        ensure_local_data_root(output_root, project_root)
        failure_dir = checked_output_path(output_root / "failures" / args.run_id, project_root)
        failure_dir.mkdir(parents=True, exist_ok=True)
        path = failure_dir / "failure-receipt.json"
        if not path.exists():
            write_json(path, {
                "schemaVersion": "manual-history-conversion-failure-1.0.0",
                "runId": args.run_id,
                "status": "FAILED_ISOLATED",
                "errorCode": error.code,
                "message": str(error),
                "databaseImportAttempted": False,
                "sourceModified": False,
            })
    except Exception:
        pass


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Convert a manual history workbook into staged, auditable daily facts")
    parser.add_argument("--source", required=True)
    parser.add_argument("--expected-source-sha256", required=True)
    parser.add_argument("--contract", required=True)
    parser.add_argument("--project-root", required=True)
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--project-ref", default="sampleproject")
    parser.add_argument("--account-id")
    parser.add_argument("--account-evidence")
    parser.add_argument("--completeness-evidence")
    parser.add_argument("--auto-overlap-root")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        receipt = convert(args)
    except ConversionError as error:
        write_failure_receipt(args, error)
        print(json.dumps({"status": "FAILED_ISOLATED", "errorCode": error.code, "message": str(error)}, ensure_ascii=False))
        return 2
    except Exception:
        error = ConversionError("UNEXPECTED_FAILURE", "Unexpected conversion failure; details intentionally omitted from the safe receipt")
        write_failure_receipt(args, error)
        traceback.print_exc(file=sys.stderr)
        print(json.dumps({"status": "FAILED_ISOLATED", "errorCode": error.code}, ensure_ascii=False))
        return 3
    summary = {
        "status": receipt["result"],
        "bundleId": receipt["bundleId"],
        "selectedFactRows": receipt["workbook"]["selectedFactRows"],
        "calendarDays": receipt["daily"]["calendarDays"],
        "daysWithRows": receipt["daily"]["daysWithRows"],
        "stagedUnverifiedDays": receipt["daily"]["stagedUnverifiedDays"],
        "overlapMatchedDays": receipt["overlap"]["matchedDays"],
        "readbackStatus": receipt["readback"]["status"],
        "outputTreeSha256": receipt["output"]["treeSha256"],
    }
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
