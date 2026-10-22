import copy
import datetime as dt
import tempfile
import unittest
from pathlib import Path

import openpyxl

from convert_manual_history import convert_workbook_rows, daily_content_hash, load_json
from verify_manual_history_bundle import VerificationError, verify_source_rows


PROJECT_ROOT = Path(__file__).resolve().parents[2]


class IndependentWorkbookVerificationTest(unittest.TestCase):
    def setUp(self):
        test_root = PROJECT_ROOT / "qa" / "runs"
        test_root.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="manual-history-source-verify-", dir=test_root)
        self.source = Path(self.temporary.name) / "synthetic.xlsx"
        self.contract = load_json(PROJECT_ROOT / "contracts/manual-history-workbook-v1.json")
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = self.contract["worksheet"]
        self.header = [f"synthetic-header-{column}" for column in range(37)]
        sheet.append(self.header)
        for index, date in enumerate((dt.date(2025, 11, 1), dt.date(2026, 1, 1), dt.date(2026, 4, 27))):
            row = [None] * 37
            row[:4] = [date, 1, f"synthetic-{index}", f"synthetic-id-{index}"]
            row[6] = dt.datetime(2025, 10, 1, 8, 30)
            if index == 0:
                row[9] = date
                row[10:29] = [column / 10 for column in range(10, 29)]
            else:
                row[9:37] = [column / 10 for column in range(9, 37)]
            if index == 2:
                row[9] = 0
                row[2] = None
            sheet.append(row)
        sheet.append([None, "排序"])
        workbook.save(self.source)
        converted, _ = convert_workbook_rows(self.source, self.contract)
        self.rows = {date.isoformat(): rows for date, rows in converted.items()}

    def tearDown(self):
        self.temporary.cleanup()

    def test_reparses_source_and_compares_all_three_mapping_periods(self):
        result = verify_source_rows(self.source, self.contract, self.rows)
        self.assertEqual(result["sourceRowsCompared"], 3)
        self.assertGreaterEqual(result["fieldValuesCompared"], 3 * 33)

    def test_detects_value_change_even_when_output_hash_is_recomputed(self):
        for field in ("stat_cost_for_roi2", "live_cvr_rate_for_roi2_v2", "video_play_duration_10s_rate_for_roi2"):
            with self.subTest(field=field):
                rows = copy.deepcopy(self.rows)
                day = rows["2025-11-01"]
                before = daily_content_hash(day)
                day[0][field] = 0 if isinstance(day[0][field], str) else day[0][field] + 1
                self.assertNotEqual(before, daily_content_hash(day))
                with self.assertRaisesRegex(VerificationError, "Source value mismatch"):
                    verify_source_rows(self.source, self.contract, rows)

    def test_detects_dropped_source_row_even_with_self_consistent_output_counts(self):
        rows = copy.deepcopy(self.rows)
        rows.pop("2026-01-01")
        with self.assertRaisesRegex(VerificationError, "Source fact missing"):
            verify_source_rows(self.source, self.contract, rows)

    def test_detects_reused_or_invented_source_row_reference(self):
        for replacement in (2, 1000):
            with self.subTest(replacement=replacement):
                rows = copy.deepcopy(self.rows)
                rows["2026-01-01"][0]["_manual_history"]["source_workbook_row"] = replacement
                with self.assertRaises(VerificationError):
                    verify_source_rows(self.source, self.contract, rows)

    def test_detects_material_row_silently_ignored_for_text_rank(self):
        workbook = openpyxl.load_workbook(self.source)
        workbook[self.contract["worksheet"]]["B3"] = "1"
        workbook.save(self.source)
        converted, _ = convert_workbook_rows(self.source, self.contract)
        rows = {date.isoformat(): items for date, items in converted.items()}
        with self.assertRaisesRegex(VerificationError, "Potential fact omitted"):
            verify_source_rows(self.source, self.contract, rows)

    def test_ignores_dated_repeated_header_with_blank_rank_and_only_metric_labels(self):
        workbook = openpyxl.load_workbook(self.source)
        header = list(self.header)
        header[0] = dt.date(2026, 5, 1)
        header[1] = None
        header[18] = "synthetic-alternate-metric-heading"
        workbook[self.contract["worksheet"]].append(header)
        workbook.save(self.source)
        result = verify_source_rows(self.source, self.contract, self.rows)
        self.assertEqual(result["sourceRowsCompared"], 3)
        self.assertEqual(result["datedRepeatedHeaderRowsSkipped"], 1)

    def test_blank_rank_numeric_fact_cannot_be_hidden_as_repeated_header(self):
        workbook = openpyxl.load_workbook(self.source)
        header = list(self.header)
        header[0] = dt.date(2026, 5, 1)
        header[1] = None
        header[10] = 12.5
        workbook[self.contract["worksheet"]].append(header)
        workbook.save(self.source)
        with self.assertRaisesRegex(VerificationError, "Potential fact omitted"):
            verify_source_rows(self.source, self.contract, self.rows)


if __name__ == "__main__":
    unittest.main()
